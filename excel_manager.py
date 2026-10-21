import os
import shutil
import time
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


MEDIA_COLUMNS: List[str] = (
    [f"Image_{i}" for i in range(1, 11)]
    + [f"Video_{i}" for i in range(1, 6)]
)

ALL_STANDARD_COLUMNS: List[str] = ["Dossier", "Statut"] + MEDIA_COLUMNS


class ExcelPermissionError(Exception):
    """Erreur levée lorsqu'un fichier Excel est ouvert dans une autre application."""
    pass


class ExcelManager:
    """
    Gère la lecture et la mise à jour du fichier Excel unique (initial).
    - Ne supprime jamais les colonnes ou informations existantes du fichier Excel.
    - Ajoute ou met à jour `Statut`, `Image_1`..`Image_10`, `Video_1`..`Video_5` dans les colonnes suivantes.
    - Colorie "Dossier vide" en rouge (#E51B23).
    - Protège le fichier contre les coupures brusques de courant (écriture atomique + copie _backup.xlsx).
    """

    @staticmethod
    def _normalize_header(val: Any) -> str:
        if val is None:
            return ""
        return str(val).strip().lower()

    @staticmethod
    def is_completed_status(val: Any) -> bool:
        """Vérifie si la valeur de la cellule Statut correspond à 'Terminé'."""
        if val is None:
            return False
        normalized = str(val).strip().lower()
        return normalized in ("terminé", "termine", "terminée", "terminee")

    @staticmethod
    def is_empty_folder_status(val: Any) -> bool:
        """Vérifie si la valeur de la cellule Statut correspond à 'Dossier vide'."""
        if val is None:
            return False
        normalized = str(val).strip().lower()
        return normalized == "dossier vide"

    @staticmethod
    def _get_backup_path(path: Path) -> Path:
        return path.with_name(f"{path.stem}_backup{path.suffix}")

    @staticmethod
    def _get_temp_path(path: Path) -> Path:
        return path.with_name(f".~tmp_{path.stem}{path.suffix}")

    @classmethod
    def _load_workbook_safe(cls, path: Path) -> openpyxl.Workbook:
        """
        Charge un classeur Excel en toute sécurité.
        Si le fichier principal a été endommagé par une coupure électrique brutale,
        restaure automatiquement la copie de secours `_backup.xlsx`.
        """
        bak_path = cls._get_backup_path(path)
        try:
            return openpyxl.load_workbook(path)
        except Exception as primary_err:
            if bak_path.exists():
                try:
                    wb = openpyxl.load_workbook(bak_path)
                    shutil.copy2(bak_path, path)
                    return wb
                except Exception:
                    pass
            raise primary_err

    @classmethod
    def _ensure_columns(cls, ws: openpyxl.worksheet.worksheet.Worksheet) -> Tuple[Dict[str, int], bool]:
        """
        Repère les colonnes existantes ('Dossier', 'Statut', 'Image_1'..'Image_10', 'Video_1'..'Video_5')
        sans jamais supprimer ni écraser les autres colonnes du fichier Excel.
        Crée dans les colonnes suivantes celles qui manquent encore.
        Retourne (mapping_colonnes, a_ete_modifie).
        """
        col_map: Dict[str, int] = {}
        max_col = max(ws.max_column or 1, 1)

        # Recherche des colonnes déjà présentes sur la ligne 1
        for col_idx in range(1, max_col + 1):
            raw_val = ws.cell(row=1, column=col_idx).value
            norm = cls._normalize_header(raw_val)
            if not norm:
                continue
            if norm in ("dossier", "dossiers", "nom_dossier", "nom du dossier") and "Dossier" not in col_map:
                col_map["Dossier"] = col_idx
            elif norm in ("statut", "status", "état", "etat") and "Statut" not in col_map:
                col_map["Statut"] = col_idx
            else:
                for media_col in MEDIA_COLUMNS:
                    if norm == media_col.lower() and media_col not in col_map:
                        col_map[media_col] = col_idx

        if "Dossier" not in col_map:
            raise ValueError(
                "La colonne 'Dossier' est introuvable dans la première ligne du fichier Excel."
            )

        # Déterminer la dernière colonne réellement utilisée sur la ligne 1
        last_used_col = 0
        for col_idx in range(1, max_col + 1):
            if ws.cell(row=1, column=col_idx).value is not None:
                last_used_col = col_idx

        modified = False
        header_fill = PatternFill(start_color="102A5E", end_color="102A5E", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)

        # Ajout des colonnes manquantes à la suite (Statut, puis Image_1..10, puis Video_1..5)
        for col_name in ["Statut"] + MEDIA_COLUMNS:
            if col_name not in col_map:
                last_used_col += 1
                col_map[col_name] = last_used_col
                c = ws.cell(row=1, column=last_used_col, value=col_name)
                c.fill = header_fill
                c.font = header_font
                c.alignment = Alignment(horizontal="center", vertical="center")
                col_letter = get_column_letter(last_used_col)
                ws.column_dimensions[col_letter].width = 20 if col_name == "Statut" else 36
                modified = True

        return col_map, modified

    @classmethod
    def _apply_status_style(cls, cell: openpyxl.cell.cell.Cell, status: str) -> None:
        """
        Applique le style visuel au statut :
        - 'Dossier vide' : colorié en ROUGE (texte rouge vif gras + fond rouge clair)
        - 'Terminé' : texte vert gras
        - 'En cours' : texte bleu gras
        """
        if cls.is_empty_folder_status(status):
            cell.font = Font(color="E51B23", bold=True)
            cell.fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
        elif cls.is_completed_status(status):
            cell.font = Font(color="16A34A", bold=True)
            cell.fill = PatternFill(fill_type=None)
        elif status.lower().startswith("en cours"):
            cell.font = Font(color="0878F9", bold=True)
            cell.fill = PatternFill(fill_type=None)
        elif status.lower().startswith("erreur"):
            cell.font = Font(color="E51B23", bold=True)
            cell.fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
        else:
            cell.font = Font(bold=False)
            cell.fill = PatternFill(fill_type=None)

    @classmethod
    def read_source_folders(cls, source_path: str) -> List[Dict[str, Any]]:
        """
        Lit le fichier Excel initial.
        S'assure que les colonnes 'Statut', 'Image_1'..'Image_10' et 'Video_1'..'Video_5' existent
        dans les colonnes suivantes sans supprimer aucune information déjà présente.
        """
        path = Path(source_path)
        bak_path = cls._get_backup_path(path)
        if not path.exists() and not bak_path.exists():
            raise FileNotFoundError(f"Le fichier Excel est introuvable : {source_path}")

        wb = cls._load_workbook_safe(path)
        ws = wb.active

        col_map, need_save = cls._ensure_columns(ws)
        dossier_col = col_map["Dossier"]
        statut_col = col_map["Statut"]

        folders: List[Dict[str, Any]] = []
        for row_idx in range(2, (ws.max_row or 1) + 1):
            raw_dossier = ws.cell(row=row_idx, column=dossier_col).value
            if raw_dossier is None or str(raw_dossier).strip() == "":
                continue
            dossier_name = str(raw_dossier).strip()
            raw_statut = ws.cell(row=row_idx, column=statut_col).value
            statut_str = "" if raw_statut is None else str(raw_statut).strip()
            completed = cls.is_completed_status(statut_str)
            empty_status = cls.is_empty_folder_status(statut_str)

            # Lecture des liens existants éventuels sur cette ligne
            existing_images: List[str] = []
            for i in range(1, 11):
                v = ws.cell(row=row_idx, column=col_map[f"Image_{i}"]).value
                existing_images.append("" if v is None else str(v).strip())

            existing_videos: List[str] = []
            for i in range(1, 6):
                v = ws.cell(row=row_idx, column=col_map[f"Video_{i}"]).value
                existing_videos.append("" if v is None else str(v).strip())

            folders.append(
                {
                    "row_index": row_idx,
                    "dossier": dossier_name,
                    "statut": statut_str,
                    "completed": completed,
                    "empty_status": empty_status,
                    "existing_images": existing_images,
                    "existing_videos": existing_videos,
                }
            )

        if need_save:
            cls._save_workbook_atomic(wb, path)
        wb.close()
        return folders

    @classmethod
    def _find_target_row(
        cls,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        dossier_col: int,
        row_index: int,
        dossier_name: str,
    ) -> int:
        current_val = ws.cell(row=row_index, column=dossier_col).value
        if current_val is not None and str(current_val).strip() == dossier_name:
            return row_index
        for r in range(2, (ws.max_row or 1) + 1):
            v = ws.cell(row=r, column=dossier_col).value
            if v is not None and str(v).strip() == dossier_name:
                return r
        return row_index

    @classmethod
    def update_source_status(
        cls,
        source_path: str,
        row_index: int,
        dossier_name: str,
        status: str = "Terminé",
    ) -> None:
        """
        Met à jour immédiatement et atomiquement la colonne 'Statut'
        pour la ligne correspondante dans le fichier Excel.
        Si status == 'Dossier vide', colorie la cellule en rouge.
        """
        path = Path(source_path)
        wb = cls._load_workbook_safe(path)
        ws = wb.active

        col_map, _ = cls._ensure_columns(ws)
        target_row = cls._find_target_row(ws, col_map["Dossier"], row_index, dossier_name)

        status_cell = ws.cell(row=target_row, column=col_map["Statut"], value=status)
        cls._apply_status_style(status_cell, status)

        cls._save_workbook_atomic(wb, path)
        wb.close()

    @classmethod
    def update_row_links(
        cls,
        excel_path: str,
        row_index: int,
        dossier_name: str,
        image_links: List[str],
        video_links: List[str],
        status: Optional[str] = None,
    ) -> None:
        """
        Met à jour uniquement les colonnes `Image_1`..`Image_10` et `Video_1`..`Video_5`
        (ainsi que `Statut` si fourni) sur la ligne du dossier dans le fichier Excel initial,
        sans jamais supprimer les autres lignes ou colonnes existantes du fichier.
        - Si un fichier photo/vidéo a été ajouté ou modifié, son lien est ajouté/mis à jour.
        - Si un fichier photo/vidéo a été supprimé du dossier source, l'ancien lien en trop est effacé.
        """
        path = Path(excel_path)
        wb = cls._load_workbook_safe(path)
        ws = wb.active

        col_map, _ = cls._ensure_columns(ws)
        target_row = cls._find_target_row(ws, col_map["Dossier"], row_index, dossier_name)

        for i in range(1, 11):
            col_idx = col_map[f"Image_{i}"]
            new_val = image_links[i - 1] if (i - 1) < len(image_links) else ""
            cell = ws.cell(row=target_row, column=col_idx)
            if (cell.value or "") != new_val:
                cell.value = new_val if new_val != "" else None

        for i in range(1, 6):
            col_idx = col_map[f"Video_{i}"]
            new_val = video_links[i - 1] if (i - 1) < len(video_links) else ""
            cell = ws.cell(row=target_row, column=col_idx)
            if (cell.value or "") != new_val:
                cell.value = new_val if new_val != "" else None

        if status is not None:
            status_cell = ws.cell(row=target_row, column=col_map["Statut"], value=status)
            cls._apply_status_style(status_cell, status)

        cls._save_workbook_atomic(wb, path)
        wb.close()

    @classmethod
    def create_source_from_directory(cls, excel_path: str, parent_dir: Optional[str] = None) -> int:
        """
        Crée ou initialise un fichier Excel `dossiers.xlsx` avec toutes les colonnes :
        Dossier | Statut | Image_1 ... Image_10 | Video_1 ... Video_5
        """
        path = Path(excel_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        subfolders: List[str] = []
        if parent_dir and os.path.isdir(parent_dir):
            try:
                subfolders = sorted(
                    [
                        d.name
                        for d in Path(parent_dir).iterdir()
                        if d.is_dir() and not d.name.startswith(".")
                    ],
                    key=lambda s: s.lower(),
                )
            except Exception:
                subfolders = []

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dossiers"

        header_fill = PatternFill(start_color="102A5E", end_color="102A5E", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)

        for col_idx, h in enumerate(ALL_STANDARD_COLUMNS, start=1):
            c = ws.cell(row=1, column=col_idx, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, folder_name in enumerate(subfolders, start=2):
            ws.cell(row=row_idx, column=1, value=folder_name)
            ws.cell(row=row_idx, column=2, value="")

        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 20
        for col_idx in range(3, len(ALL_STANDARD_COLUMNS) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 36
        ws.row_dimensions[1].height = 24
        ws.freeze_panes = "C2"

        cls._save_workbook_atomic(wb, path)
        wb.close()
        return len(subfolders)

    @classmethod
    def _save_workbook_atomic(
        cls,
        wb: openpyxl.Workbook,
        path: Path,
        retries: int = 6,
        delay: float = 0.8,
    ) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp_path = cls._get_temp_path(path)
        bak_path = cls._get_backup_path(path)

        last_err: Optional[Exception] = None
        for attempt in range(1, retries + 1):
            try:
                wb.save(tmp_path)

                if path.exists() and path.stat().st_size > 0:
                    try:
                        shutil.copy2(path, bak_path)
                    except Exception:
                        pass

                os.replace(tmp_path, path)

                try:
                    shutil.copy2(path, bak_path)
                except Exception:
                    pass
                return

            except PermissionError as e:
                last_err = e
                time.sleep(delay * attempt)
            except Exception as e:
                last_err = e
                time.sleep(delay)
            finally:
                if tmp_path.exists():
                    try:
                        tmp_path.unlink()
                    except Exception:
                        pass

        if isinstance(last_err, PermissionError):
            raise ExcelPermissionError(
                f"Impossible d'enregistrer '{path.name}'. Veuillez fermer ce fichier dans Microsoft Excel puis réessayer."
            ) from last_err
        if last_err:
            raise last_err
