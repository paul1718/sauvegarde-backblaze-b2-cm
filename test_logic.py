import os
import tempfile
from pathlib import Path
from typing import Dict, Any, List, Tuple
import openpyxl
import pandas as pd

from config_manager import ConfigManager
from excel_manager import ExcelManager, ALL_STANDARD_COLUMNS
from b2_uploader import (
    B2UploaderService,
    resolve_matching_folder,
    is_tav_folder,
)
from generate_icon import create_app_icon


class FakeB2FileVersion:
    def __init__(self, file_name: str, size: int, file_id: str) -> None:
        self.file_name = file_name
        self.size = size
        self.id_ = file_id


class FakeB2Bucket:
    def __init__(self) -> None:
        self.files: Dict[str, FakeB2FileVersion] = {}
        self.upload_calls: List[str] = []
        self.delete_calls: List[str] = []

    def ls(self, folder_to_list: str = "", recursive: bool = True) -> List[Tuple[FakeB2FileVersion, str]]:
        return [
            (fv, "")
            for name, fv in list(self.files.items())
            if name.startswith(folder_to_list)
        ]

    def upload_local_file(self, local_file: str, file_name: str) -> FakeB2FileVersion:
        sz = Path(local_file).stat().st_size
        fv = FakeB2FileVersion(file_name, sz, f"id_{len(self.files) + 1}")
        self.files[file_name] = fv
        self.upload_calls.append(file_name)
        return fv

    def delete_file_version(self, file_id: str, file_name: str) -> None:
        self.files.pop(file_name, None)
        self.delete_calls.append(file_name)

    def get_download_url(self, file_name: str) -> str:
        return f"https://f000.backblazeb2.com/file/mon-bucket/{file_name}"


def run_tests() -> None:
    create_app_icon()

    print("1. Test du fichier Excel unique déjà rempli (conservation des données existantes + colonnes suivantes)...")
    with tempfile.TemporaryDirectory() as tmpdir:
        parent = Path(tmpdir)
        folder_prefix = parent / "TAV_204 - Terrain Bastos (Lot 3)"
        folder_empty_tav = parent / "TAV_999 Parcelle_Vide"
        folder_empty_other = parent / "Anniversaire_Vide_2026"

        folder_prefix.mkdir()
        folder_empty_tav.mkdir()
        folder_empty_other.mkdir()

        # Création d'un fichier Excel initial déjà rempli avec une colonne métier "Quartier" et des liens existants
        excel_path = parent / "dossiers.xlsx"
        df_init = pd.DataFrame(
            [
                {
                    "Dossier": "Dossier_Autre_Disque",
                    "Quartier": "Odza",
                    "Statut": "Terminé",
                    "Image_1": "https://f000.backblazeb2.com/file/mon-bucket/Dossier_Autre_Disque/ancienne_img.jpg",
                },
                {
                    "Dossier": "TAV_204",
                    "Quartier": "Bastos",
                    "Statut": "Terminé",
                    "Image_1": "https://f000.backblazeb2.com/file/mon-bucket/old1.jpg",
                    "Image_2": "https://f000.backblazeb2.com/file/mon-bucket/old2_supprimee.jpg",
                },
                {
                    "Dossier": "TAV_999",
                    "Quartier": "Mendong",
                    "Statut": "",
                    "Image_1": "",
                },
                {
                    "Dossier": "Anniversaire",
                    "Quartier": "Biyem-Assi",
                    "Statut": "",
                    "Image_1": "",
                },
            ]
        )
        df_init.to_excel(excel_path, index=False)

        folders = ExcelManager.read_source_folders(str(excel_path))
        assert len(folders) == 4

        print("2. Test de la mise à jour des liens (modification / suppression de lien selon l'état actuel du dossier source)...")
        # Dans TAV_204, il n'y a désormais qu'une seule photo (photo_actuelle.jpg) et 1 vidéo (clip.mp4) :
        # L'ancien lien Image_2 doit donc être supprimé et Image_1 mis à jour !
        (folder_prefix / "photo_actuelle.jpg").write_bytes(b"new_img_bytes")
        (folder_prefix / "clip.mp4").write_bytes(b"new_vid_bytes")

        uploader = B2UploaderService("k", "a", "b")
        uploader.bucket = FakeB2Bucket()

        imgs, vids, _, _, _ = uploader.upload_folder(folder_prefix)
        ExcelManager.update_row_links(
            excel_path=str(excel_path),
            row_index=folders[1]["row_index"],
            dossier_name="TAV_204",
            image_links=imgs,
            video_links=vids,
            status="Terminé",
        )

        print("3. Test des dossiers vides et de la coloration en ROUGE de 'Dossier vide'...")
        std_img_file = parent / "image_standard_unique.jpg"
        std_img_file.write_bytes(b"standard_image_bytes_12345")
        std_url = uploader.ensure_tav_standard_image_url(str(std_img_file))

        # Dossier TAV_999 vide -> Image standard + "Dossier vide" en rouge
        ExcelManager.update_row_links(
            excel_path=str(excel_path),
            row_index=folders[2]["row_index"],
            dossier_name="TAV_999",
            image_links=[std_url],
            video_links=[],
            status="Dossier vide",
        )

        # Dossier Anniversaire vide (non TAV) -> Aucun lien + "Dossier vide" en rouge
        ExcelManager.update_row_links(
            excel_path=str(excel_path),
            row_index=folders[3]["row_index"],
            dossier_name="Anniversaire",
            image_links=[],
            video_links=[],
            status="Dossier vide",
        )

        # Vérification complète du classeur Excel finalisé
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        col_map, _ = ExcelManager._ensure_columns(ws)

        # 1) Vérifier que la ligne 2 (Dossier_Autre_Disque) et la colonne "Quartier" n'ont PAS été supprimées ni écrasées
        assert ws.cell(row=2, column=2).value == "Odza", "La colonne personnalisée Quartier doit être intacte !"
        assert (
            ws.cell(row=2, column=col_map["Image_1"]).value
            == "https://f000.backblazeb2.com/file/mon-bucket/Dossier_Autre_Disque/ancienne_img.jpg"
        ), "Les liens existants des dossiers non touchés doivent être préservés !"

        # 2) Vérifier que TAV_204 (ligne 3) a bien mis à jour Image_1, supprimé l'ancien Image_2, et ajouté Video_1
        assert ws.cell(row=3, column=col_map["Image_1"]).value == imgs[0]
        assert ws.cell(row=3, column=col_map["Image_2"]).value is None, "L'ancien lien Image_2 supprimé du dossier source doit être effacé !"
        assert ws.cell(row=3, column=col_map["Video_1"]).value == vids[0]

        # 3) Vérifier que TAV_999 (ligne 4) et Anniversaire (ligne 5) ont "Dossier vide" colorié en ROUGE (#E51B23)
        cell_statut_tav = ws.cell(row=4, column=col_map["Statut"])
        cell_statut_other = ws.cell(row=5, column=col_map["Statut"])
        assert cell_statut_tav.value == "Dossier vide"
        assert "E51B23" in str(cell_statut_tav.font.color.rgb), f"Couleur obtenue : {cell_statut_tav.font.color.rgb}"
        assert cell_statut_other.value == "Dossier vide"
        assert "E51B23" in str(cell_statut_other.font.color.rgb)
        assert ws.cell(row=4, column=col_map["Image_1"]).value == std_url
        assert ws.cell(row=5, column=col_map["Image_1"]).value is None
        wb.close()
        print("   -> OK : Fichier Excel unique, préservation des données, mise à jour/suppression de liens et 'Dossier vide' en rouge validés !")

        print("4. Test de simulation de coupure de courant brutale (fichier Excel tronqué + journal WAL en attente)...")
        # Simuler une écriture interrompue par une coupure électrique :
        # 1) On écrit une transaction dans le journal WAL
        ExcelManager._record_wal_entry(
            path=excel_path,
            row_index=folders[3]["row_index"],
            dossier_name="Anniversaire",
            image_links=["https://f000.backblazeb2.com/file/mon-bucket/Anniversaire/photo_secours.jpg"],
            video_links=[],
            status="Terminé",
        )
        # 2) On détruit/tronque le fichier Excel principal comme si le courant avait coupé en plein milieu d'un accès disque
        excel_path.write_bytes(b"\x00\x00COURANT_COUPE_BRUTALEMENT_FICHIER_TRONQUE")
        assert not ExcelManager._is_valid_xlsx_file(excel_path), "Le fichier tronqué doit être détecté comme invalide !"

        # 3) À la réouverture (ou à l'opération suivante), ExcelManager doit auto-réparer le fichier depuis le backup/miroir ET rejouer le WAL !
        recovered_folders = ExcelManager.read_source_folders(str(excel_path))
        assert ExcelManager._is_valid_xlsx_file(excel_path), "Le fichier Excel principal doit être intégralement restauré !"
        assert not ExcelManager._get_wal_path(excel_path).exists(), "Le journal WAL doit avoir été rejoué et nettoyé !"

        wb_rec = openpyxl.load_workbook(excel_path)
        ws_rec = wb_rec.active
        col_map_rec, _ = ExcelManager._ensure_columns(ws_rec)
        assert ws_rec.cell(row=5, column=col_map_rec["Statut"]).value == "Terminé"
        assert (
            ws_rec.cell(row=5, column=col_map_rec["Image_1"]).value
            == "https://f000.backblazeb2.com/file/mon-bucket/Anniversaire/photo_secours.jpg"
        )
        wb_rec.close()
        print("   -> OK : Auto-réparation après coupure de courant + rejeu automatique du journal WAL validés à 100% !")

        print("5. Test de l'enregistrement du lancement automatique Windows au rallumage du PC après coupure de courant...")
        cm = ConfigManager()
        assert cm.ensure_windows_run_registry(), "L'enregistrement dans le Registre Windows (HKCU Run) doit réussir !"
        cm.mark_session_active(True, current_folder="TAV_204", last_completed_row=3)
        state_active = cm.load_state_journal()
        assert state_active["active_session"] is True
        assert state_active["current_folder"] == "TAV_204"
        assert state_active["last_completed_row_in_run"] == 3
        vbs_startup = cm._get_startup_vbs_path()
        if vbs_startup is not None:
            assert vbs_startup.exists(), "Le lanceur de secours dans le dossier Démarrage Windows doit être présent pendant une session active !"
        cm.mark_session_active(False)
        if vbs_startup is not None:
            assert not vbs_startup.exists(), "Le lanceur de secours doit être nettoyé une fois la sauvegarde terminée !"
        print("   -> OK : Lancement automatique au rallumage du PC et reprise exacte au point d'arrêt validés !")

        print("6. Test de la résolution par préfixe avec espace obligatoire (ex: 'TV326' ne doit JAMAIS correspondre à 'TV3260')...")
        test_dir = parent / "test_prefixes"
        test_dir.mkdir()
        dir_tv3260 = test_dir / "TV3260"
        dir_tv3260_desc = test_dir / "TV3260 - Terrain B"
        dir_tv326_desc = test_dir / "TV326 - Terrain A"
        dir_tv326_exact = test_dir / "TV326"

        dir_tv3260.mkdir()
        dir_tv3260_desc.mkdir()

        # Si TV3260 et TV3260 - Terrain B sont seuls, TV326 ne doit RIEN matcher (pas de fausse association)
        assert resolve_matching_folder(test_dir, "TV326") is None, (
            "TV326 ne doit JAMAIS correspondre à TV3260 car il manque l'espace après le préfixe !"
        )

        # Dès qu'un dossier avec espace existe ('TV326 - Terrain A'), TV326 doit l'associer
        dir_tv326_desc.mkdir()
        res_prefix = resolve_matching_folder(test_dir, "TV326")
        assert res_prefix is not None and res_prefix.name == "TV326 - Terrain A"

        # Si le dossier exact 'TV326' existe aussi, il doit être prioritaire
        dir_tv326_exact.mkdir()
        res_exact = resolve_matching_folder(test_dir, "TV326")
        assert res_exact is not None and res_exact.name == "TV326"

        # Et TV3260 doit bien correspondre à TV3260
        res_tv3260 = resolve_matching_folder(test_dir, "TV3260")
        assert res_tv3260 is not None and res_tv3260.name == "TV3260"

        # Tests exhaustifs sur les exemples réels fournis par l'utilisateur :
        # Ex: TV42 ne doit JAMAIS correspondre à 'TV420 pas de photo'
        # Ex: TV155 ne doit JAMAIS correspondre à 'TV1550'
        # Ex: TV42 -> 'TV42' ou 'TV42 niofiofg'
        # Ex: TV420 -> 'TV420 pas de photo'
        # Ex: TV4 -> 'TV4' ou 'TV4 pas de photo'
        # Ex: TV4250 -> 'TV4250 pas de photo'
        # Ex: TV42502030 -> 'TV42502030 pas de photo'
        # Ex: TV42505948 -> 'TV42505948'
        user_test_dir = parent / "user_test_prefixes"
        user_test_dir.mkdir()
        for name in [
            "TV420 pas de photo",
            "TV1550",
            "TV42 niofiofg",
            "TV4 pas de photo",
            "TV4250 pas de photo",
            "TV42502030 pas de photo",
            "TV42505948",
        ]:
            (user_test_dir / name).mkdir()

        # TV42 ne doit PAS matcher TV420 pas de photo, mais TV42 niofiofg
        res_tv42 = resolve_matching_folder(user_test_dir, "TV42")
        assert res_tv42 is not None and res_tv42.name == "TV42 niofiofg"

        # TV155 ne doit PAS matcher TV1550 (manque espace)
        assert resolve_matching_folder(user_test_dir, "TV155") is None

        # TV420 doit matcher TV420 pas de photo
        res_tv420 = resolve_matching_folder(user_test_dir, "TV420")
        assert res_tv420 is not None and res_tv420.name == "TV420 pas de photo"

        # TV4 doit matcher TV4 pas de photo
        res_tv4 = resolve_matching_folder(user_test_dir, "TV4")
        assert res_tv4 is not None and res_tv4.name == "TV4 pas de photo"

        # Si on ajoute TV4 exact, il devient prioritaire sur TV4 pas de photo
        (user_test_dir / "TV4").mkdir()
        res_tv4_exact = resolve_matching_folder(user_test_dir, "TV4")
        assert res_tv4_exact is not None and res_tv4_exact.name == "TV4"

        # TV4250 doit matcher TV4250 pas de photo
        res_tv4250 = resolve_matching_folder(user_test_dir, "TV4250")
        assert res_tv4250 is not None and res_tv4250.name == "TV4250 pas de photo"

        # TV42502030 doit matcher TV42502030 pas de photo
        res_tv_long = resolve_matching_folder(user_test_dir, "TV42502030")
        assert res_tv_long is not None and res_tv_long.name == "TV42502030 pas de photo"

        # TV42505948 exact
        res_tv_exact_num = resolve_matching_folder(user_test_dir, "TV42505948")
        assert res_tv_exact_num is not None and res_tv_exact_num.name == "TV42505948"

        print("   -> OK : Règle de l'espace obligatoire après chaque préfixe validée à 100% sur TOUS les exemples !")

    # Mise à jour du modèle dossiers.xlsx du projet avec les colonnes complètes s'il est vide
    proj_excel = Path(__file__).resolve().parent / "dossiers.xlsx"
    if proj_excel.exists():
        ExcelManager.read_source_folders(str(proj_excel))

    print("\nTOUS LES TESTS SONT PASSÉS AVEC SUCCÈS !")


if __name__ == "__main__":
    run_tests()
